import sys
from pptx import Presentation
import json

def extract_text(pptx_file):
    prs = Presentation(pptx_file)
    extracted = []
    
    for s_idx, slide in enumerate(prs.slides):
        for sh_idx, shape in enumerate(slide.shapes):
            if hasattr(shape, "text_frame") and shape.text_frame:
                for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                    text = paragraph.text.strip()
                    if text:
                        extracted.append({
                            "s_idx": s_idx,
                            "sh_idx": sh_idx,
                            "p_idx": p_idx,
                            "text": text,
                            "is_cell": False
                        })
            elif hasattr(shape, "has_table") and shape.has_table:
                for r_idx, row in enumerate(shape.table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        if hasattr(cell, "text_frame") and cell.text_frame:
                            for p_idx, paragraph in enumerate(cell.text_frame.paragraphs):
                                text = paragraph.text.strip()
                                if text:
                                    extracted.append({
                                        "s_idx": s_idx,
                                        "sh_idx": sh_idx,
                                        "r_idx": r_idx,
                                        "c_idx": c_idx,
                                        "p_idx": p_idx,
                                        "text": text,
                                        "is_cell": True
                                    })
    
    with open("pptx_text.json", "w", encoding="utf-8") as f:
        json.dump(extracted, f, ensure_ascii=False, indent=2)

if __name__ == "__main__":
    extract_text("X-AURUM_Presentation.pptx")
